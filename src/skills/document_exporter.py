"""
skills/document_exporter.py — Kết xuất bảng số liệu & báo cáo chi tiết sang Excel (.xlsx) / Word (.docx).
Giải quyết vấn đề bảng Markdown vỡ dòng, co cụm, lỗi font trên Telegram di động và giúp
phản hồi voice chỉ cần đọc câu trọng tâm ngắn gọn.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from core.logger import logger
from config import REPORTS_OUTPUT_DIR

Path(REPORTS_OUTPUT_DIR).mkdir(parents=True, exist_ok=True)


# ── Làm sạch markdown trong cell ─────────────────────────────────────────────
def clean_markdown_cell(text: str) -> str:
    """Loại bỏ markdown formatting để cell văn bản trong Excel/Word hiển thị sạch sẽ."""
    t = text.strip()
    t = re.sub(r'\*{1,3}([^*]+)\*{1,3}', r'\1', t)   # bold/italic
    t = re.sub(r'`([^`]+)`', r'\1', t)               # inline code
    t = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', t)   # link markdown -> text
    return t.strip()


# ── Trích xuất bảng Markdown ────────────────────────────────────────────────
def extract_markdown_tables(text: str) -> list[dict]:
    """
    Tìm và trích xuất tất cả bảng Markdown có trong văn bản:
    Dạng:
    | Cột 1 | Cột 2 |
    |---|---|
    | Dữ liệu 1 | Dữ liệu 2 |
    """
    tables = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("|") and line.endswith("|") and "|" in line[1:-1]:
            # Có thể là bắt đầu bảng
            table_lines = [line]
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith("|") and lines[j].strip().endswith("|"):
                table_lines.append(lines[j].strip())
                j += 1
            
            if len(table_lines) >= 3:
                # Kiểm tra dòng thứ 2 có phải separator |---| không
                sep_line = table_lines[1]
                if re.match(r'^\|(\s*:?-+:?\s*\|)+$', sep_line):
                    header = [clean_markdown_cell(c) for c in table_lines[0].strip('|').split('|')]
                    rows = []
                    for row_str in table_lines[2:]:
                        cells = [clean_markdown_cell(c) for c in row_str.strip('|').split('|')]
                        # Đảm bảo số lượng cột khớp header
                        if len(cells) < len(header):
                            cells.extend([""] * (len(header) - len(cells)))
                        elif len(cells) > len(header):
                            cells = cells[:len(header)]
                        rows.append(cells)
                    
                    if header and rows:
                        tables.append({
                            "headers": header,
                            "rows": rows,
                            "raw": "\n".join(table_lines)
                        })
            i = j
        else:
            i += 1
    return tables


def has_tabular_or_detailed_content(text: str) -> dict:
    """Kiểm tra văn bản có bảng Markdown hoặc danh sách so sánh chi tiết cần xuất file không."""
    tables = extract_markdown_tables(text)
    has_table = len(tables) > 0
    
    # Đoạn văn chi tiết: độ dài hoặc nhiều mục / lịch trình / danh sách
    text_lower = text.lower()
    has_keywords = any(kw in text_lower for kw in ["lịch trình", "kế hoạch", "hành trình", "so sánh", "danh sách"])
    has_sections = (
        text.count("\n### ") >= 2 or 
        text.count("\n## ") >= 2 or 
        text.count("\n- ") >= 4 or 
        text.count("\n* ") >= 4 or 
        text.count("\n1. ") >= 2
    )
    is_detailed = (len(text) >= 400 and (has_sections or has_keywords)) or len(text) > 650
    return {
        "has_table": has_table,
        "is_detailed": is_detailed,
        "tables": tables,
    }


# ── Tách câu đầu trọng tâm và nội dung chi tiết ─────────────────────────────
def split_core_and_detail(text: str) -> tuple[str, str]:
    """
    Tách câu đầu trọng tâm (core focus) và phần nội dung chi tiết (detail).
    Dùng câu trọng tâm để:
    1. Trả lời ngay đầu tin nhắn Telegram.
    2. Đọc nhanh qua voice (chỉ tốn vài giây).
    """
    cleaned = text.strip()
    paragraphs = [p.strip() for p in cleaned.split("\n\n") if p.strip()]
    if not paragraphs:
        return cleaned, ""
    
    # Nếu đoạn đầu là chào hỏi hoặc ngắn, có thể lấy 1-2 câu đầu
    first_para = paragraphs[0]
    
    # Nếu đoạn đầu là bảng luôn (hiếm gặp), lấy tóm tắt ngắn
    if first_para.startswith("|"):
        return "Dưới đây là bảng số liệu và phân tích so sánh chi tiết theo yêu cầu của bạn:", cleaned
    
    # Lấy 1-3 câu đầu tiên của đoạn đầu làm core_summary
    sentences = re.split(r'(?<=[.!?])\s+', first_para)
    if len(sentences) <= 3:
        core_summary = first_para
    else:
        core_summary = " ".join(sentences[:3])
    
    # Phần còn lại
    detail = "\n\n".join(paragraphs[1:]) if len(paragraphs) > 1 else ""
    return core_summary, detail


# ── Xuất Excel (.xlsx) bằng openpyxl ─────────────────────────────────────────
def create_excel_document(
    title: str,
    tables: list[dict],
    summary: str = "",
    filename_prefix: str = "bang_du_lieu",
) -> str:
    """Tạo tệp Excel định dạng đẹp mắt từ bảng Markdown."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu chi tiết"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    
    meta_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    summary_font = Font(name="Segoe UI", size=11, bold=False, color="1F1F1F")

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

    row_even_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    row_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    data_font = Font(name="Segoe UI", size=10, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'),
    )

    current_row = 1

    # Dòng 1: Banner Tiêu đề
    max_cols = max((len(t.get("headers", [])) for t in tables), default=4)
    max_cols = max(max_cols, 4)

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell = ws.cell(row=current_row, column=1, value=f"  {title.upper()}")
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[current_row].height = 36
    current_row += 1

    # Dòng 2: Thời gian xuất
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cell_meta = ws.cell(row=current_row, column=1, value=f"Thời gian tạo: {now_str} • Trợ lý AI Telegram")
    cell_meta.font = meta_font
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # Dòng 3: Tóm tắt nếu có
    if summary:
        current_row += 1
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
        c_sum = ws.cell(row=current_row, column=1, value=f"📌 Trọng tâm: {summary}")
        c_sum.font = summary_font
        c_sum.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[current_row].height = 32
        current_row += 1

    current_row += 1

    # Kết xuất từng bảng
    for t_idx, t in enumerate(tables, 1):
        headers = t.get("headers", [])
        rows = t.get("rows", [])

        # Header
        for col_idx, h in enumerate(headers, 1):
            h_cell = ws.cell(row=current_row, column=col_idx, value=h)
            h_cell.font = header_font
            h_cell.fill = header_fill
            h_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            h_cell.border = thin_border
        ws.row_dimensions[current_row].height = 28
        current_row += 1

        # Data Rows
        for r_idx, row in enumerate(rows):
            fill = row_even_fill if r_idx % 2 == 0 else row_odd_fill
            for col_idx, val in enumerate(row, 1):
                d_cell = ws.cell(row=current_row, column=col_idx, value=val)
                d_cell.font = data_font
                d_cell.fill = fill
                d_cell.border = thin_border
                d_cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[current_row].height = 24
            current_row += 1
        
        current_row += 2  # Khoảng cách giữa các bảng

    # Tự động căn chỉnh độ rộng cột
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Bỏ qua merged banner khi tính chiều rộng
            if cell.row in (1, 2) and col_letter != "A":
                continue
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

    # Lưu tệp
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    clean_prefix = re.sub(r'[^a-zA-Z0-9_]', '_', filename_prefix)[:30]
    out_path = Path(REPORTS_OUTPUT_DIR) / f"{clean_prefix}_{ts}.xlsx"
    wb.save(str(out_path))
    logger.info(f"📊 Đã tạo tệp Excel: {out_path}")
    return str(out_path)


# ── Xuất Word (.docx) bằng python-docx ───────────────────────────────────────
def create_word_document(
    title: str,
    summary: str,
    full_text: str,
    tables: list[dict],
    filename_prefix: str = "thong_tin_chi_tiet",
) -> str:
    """Tạo tệp Word (.docx) định dạng chuyên nghiệp từ văn bản và bảng số liệu."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()

    # Thiết lập lề 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Tiêu đề
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_p.paragraph_format.space_after = Pt(4)

    # Thời gian tạo
    time_p = doc.add_paragraph()
    time_run = time_p.add_run(f"Thời gian kết xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')} • Trợ lý AI")
    time_run.font.name = "Arial"
    time_run.font.size = Pt(9.5)
    time_run.font.italic = True
    time_run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    time_p.paragraph_format.space_after = Pt(14)

    # Hộp tóm tắt trọng tâm (Callout Box)
    if summary:
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_table.autofit = False
        cell = box_table.cell(0, 0)
        cell.width = Inches(6.8)

        # Đổi màu nền cell xám xanh nhạt #F0F4F8
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F0F4F8')
        tcPr.append(shd)

        cp = cell.paragraphs[0]
        c_bold = cp.add_run("📌 TÓM TẮT TRỌNG TÂM: ")
        c_bold.font.bold = True
        c_bold.font.name = "Arial"
        c_bold.font.size = Pt(10.5)
        c_bold.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        c_text = cp.add_run(summary)
        c_text.font.name = "Arial"
        c_text.font.size = Pt(10.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Loại bỏ bảng raw khỏi văn bản để tránh in 2 lần
    cleaned_text = full_text
    for t in tables:
        if "raw" in t:
            cleaned_text = cleaned_text.replace(t["raw"], "")

    # Thêm các đoạn văn bản
    paragraphs = [p.strip() for p in cleaned_text.split("\n\n") if p.strip()]
    for para in paragraphs:
        if para.startswith("### "):
            hp = doc.add_paragraph()
            hrun = hp.add_run(para.replace("### ", "").strip())
            hrun.font.name = "Arial"
            hrun.font.size = Pt(13)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
            hp.paragraph_format.space_before = Pt(10)
            hp.paragraph_format.space_after = Pt(4)
        elif para.startswith("## "):
            hp = doc.add_paragraph()
            hrun = hp.add_run(para.replace("## ", "").strip())
            hrun.font.name = "Arial"
            hrun.font.size = Pt(14)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            hp.paragraph_format.space_before = Pt(12)
            hp.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(clean_markdown_cell(para))
            run.font.name = "Arial"
            run.font.size = Pt(10.5)

    # Nhúng các bảng dữ liệu
    for t_idx, t in enumerate(tables, 1):
        headers = t.get("headers", [])
        rows = t.get("rows", [])
        if not headers or not rows:
            continue

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        tbl_p = doc.add_paragraph()
        t_label = tbl_p.add_run(f"📊 Bảng {t_idx}: Dữ liệu chi tiết")
        t_label.font.name = "Arial"
        t_label.font.bold = True
        t_label.font.size = Pt(11)
        t_label.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Định dạng Header
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            # Nền xanh đậm
            tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), '2F5597')
            tcPr.append(shd)

            for p in hdr_cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(10)

        # Định dạng Data Rows
        for r_idx, row_data in enumerate(rows, 1):
            row_cells = table.rows[r_idx].cells
            fill_hex = "F2F5F9" if r_idx % 2 == 0 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = val
                tcPr = row_cells[c_idx]._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), fill_hex)
                tcPr.append(shd)

                for p in row_cells[c_idx].paragraphs:
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Lưu tệp
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    clean_prefix = re.sub(r'[^a-zA-Z0-9_]', '_', filename_prefix)[:30]
    out_path = Path(REPORTS_OUTPUT_DIR) / f"{clean_prefix}_{ts}.docx"
    doc.save(str(out_path))
    logger.info(f"📄 Đã tạo tệp Word: {out_path}")
    return str(out_path)


# ── Hàm xuất thông minh tự chọn Excel hoặc Word ─────────────────────────────
def export_document_smart(
    query: str,
    full_text: str,
    summary: str = "",
) -> Optional[tuple[str, str]]:
    """
    Tự động phân tích và kết xuất tệp Word hoặc Excel tùy thuộc vào tính chất nội dung:
    - Nếu nội dung có bảng số liệu/so sánh: Ưu tiên Excel (.xlsx) để người dùng xem bảng cực chuẩn.
    - Nếu nội dung là bài viết phân tích, báo cáo nhiều phần: Ưu tiên Word (.docx).
    Trả về (file_path, file_type) hoặc None nếu không cần xuất.
    """
    check = has_tabular_or_detailed_content(full_text)
    tables = check["tables"]

    # Rút gọn query thành title
    title = query.strip() or "Báo cáo chi tiết AI"
    title = title.split("\n")[0][:80]

    # Prefix tên tệp an toàn
    prefix = re.sub(r'[^\w\s-]', '', query).strip().replace(" ", "_")[:20] or "Bao_cao_AI"

    if check["has_table"]:
        # Có bảng số liệu: xuất Excel để không bị lỗi cột và lỗi font
        try:
            excel_path = create_excel_document(title, tables, summary=summary, filename_prefix=prefix)
            return excel_path, "Excel"
        except Exception as e:
            logger.warning(f"⚠️ Lỗi xuất Excel: {e}")
            try:
                word_path = create_word_document(title, summary, full_text, tables, filename_prefix=prefix)
                return word_path, "Word"
            except Exception as e2:
                logger.error(f"❌ Lỗi xuất Word fallback: {e2}")
                return None
    elif check["is_detailed"]:
        # Văn bản chi tiết / lịch trình / kế hoạch: xuất Word để người dùng đọc tiện lợi và thẩm mỹ
        try:
            word_path = create_word_document(title, summary, full_text, tables, filename_prefix=prefix)
            return word_path, "Word"
        except Exception as e:
            logger.error(f"❌ Lỗi xuất Word: {e}")
            return None

    return None
